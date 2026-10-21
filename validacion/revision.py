"""Validación del módulo de revisión de Bases con los expedientes de demostración.

Cada convocatoria de demo/ trae fallas sembradas y la lista de verificaciones que deberían fallar en cada ronda
(demo/expediente.py). El script sube los documentos, ejecuta las dos rondas y compara lo detectado con lo
esperado: exhaustividad (fallas sembradas detectadas), precisión (hallazgos que eran fallas reales) y falsos
positivos por verificación.

Como evidencia, cada convocatoria queda en Supabase como un concurso «VALIDACIÓN · …» en revisión de Bases,
con sus documentos, hallazgos y pasos del agente, para verlo en el portal del comité. El reporte se guarda en
validacion/resultados/ en Markdown, JSON y Word (con los documentos usados y los resultados).

Uso:
    python -m validacion.revision                          # las 5 convocatorias, con el agente revisor
    python -m validacion.revision --sin-agente             # solo reglas fijas (no gasta Gemini en el agente)
    python -m validacion.revision --convocatorias convocatoria_4_supervision_canal_de_riego
    python -m validacion.revision --en-memoria             # no escribe en Supabase (no deja evidencia)
"""
import argparse
import json
import time
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from docx import Document

import db.client
from config import LLM_MODEL
from demo.datos import CONVOCATORIAS
from demo.expediente import EXPEDIENTES
from revision.catalogo import AGENTE, POR_CODIGO, REGLA, TIPOS_DOCUMENTO

CARPETA_DEMO = Path(__file__).parent.parent / "demo" / "convocatorias"
CARPETA_RESULTADOS = Path(__file__).parent / "resultados"
ARCHIVOS = {
    "A": {"requerimiento": "A1 Requerimiento.pdf", "certificacion": "A2 Certificación presupuestal.pdf",
          "estrategia": "A3 Estrategia de contratación.pdf", "interaccion_mercado": "A4 Interacción con el mercado.pdf",
          "designacion_evaluadores": "A5 Designación del comité.pdf",
          "aprobacion_expediente": "A6 Aprobación del expediente.pdf", "bases": "A7 Bases para convocar.pdf"},
    "B": {"pliego": "B1 Pliego de absolución.pdf"},
}


def _base_en_memoria():
    from tests.fakes import SupabaseFalso

    falsa = SupabaseFalso()
    db.client.create_client = lambda *_: falsa
    db.client.get_client.cache_clear()
    return falsa


def nombre_del_concurso(carpeta: str) -> str:
    convocatoria = next(c for c in CONVOCATORIAS if c.carpeta == carpeta)
    tema = carpeta.split("_", 2)[-1].replace("_", " ")
    return f"VALIDACIÓN · {convocatoria.nomenclatura} · {tema}"


def revisar_convocatoria(carpeta: str, con_agente: bool, en_memoria: bool = False) -> list[dict]:
    from revision import servicio

    client = _base_en_memoria() if en_memoria else db.client.get_client()
    expediente = EXPEDIENTES[carpeta]
    concurso = client.table("concursos").insert({
        "nombre": nombre_del_concurso(carpeta), "estado": "revision_bases", "ronda_revision": "A",
        "descripcion": f"Validación del módulo de revisión de Bases ({datetime.now():%d/%m/%Y}). Fallas sembradas: "
                       f"{expediente.descripcion} Hallazgos esperados: ronda A {', '.join(expediente.esperado_a) or 'ninguno'}; "
                       f"ronda B {', '.join(expediente.esperado_b) or 'ninguno'}.",
    }).execute().data[0]
    origen = CARPETA_DEMO / carpeta
    filas = []
    for ronda in ("A", "B"):
        inicio, desde = time.monotonic(), datetime.now(timezone.utc).isoformat()
        archivos = [(tipo, origen / "revision_de_bases" / nombre) for tipo, nombre in ARCHIVOS[ronda].items()]
        if ronda == "B":
            archivos.append(("bases_integradas", next(origen.glob("Bases Integradas *.pdf"))))
        documentos = []
        for tipo, ruta in archivos:
            fila, _ = servicio.subir_documento(concurso["id"], tipo, ruta.name, ruta.read_bytes())
            documentos.append({"tipo": tipo, "archivo": ruta.name, "version": fila["version"],
                               "paginas": len(fila.get("paginas") or []),
                               "ruta": str(ruta.relative_to(CARPETA_DEMO.parent.parent))})
        revision = servicio.revisar(concurso["id"], ronda, con_agente=con_agente)
        consultas = (client.table("consultas_ia").select("tokens_entrada,tokens_salida")
                     .eq("concurso_id", concurso["id"]).gte("created_at", desde).execute().data)
        filas.append({
            "convocatoria": carpeta, "concurso": concurso["nombre"], "concurso_id": concurso["id"], "ronda": ronda,
            "segundos": round(time.monotonic() - inicio, 1), "documentos": documentos,
            "esperadas": sorted(getattr(expediente, f"esperado_{ronda.lower()}")),
            "resultados": revision["resultados"], "consultas": len(consultas),
            "tokens": sum((c.get("tokens_entrada") or 0) + (c.get("tokens_salida") or 0) for c in consultas),
        })
        if ronda == "A":
            # Como al cerrar la ronda A en el portal: el concurso pasa a la ronda B. Los hallazgos de la ronda A
            # quedan sin responder, para que el comité los vea en el informe tal como los dejó el sistema.
            client.table("concursos").update({"ronda_revision": "B"}).eq("id", concurso["id"]).execute()
    return filas


def metricas(filas: list[dict], con_agente: bool) -> dict:
    evaluables = {codigo for codigo, v in POR_CODIGO.items() if con_agente or v.metodo == REGLA}
    vp = fn = fp = 0
    por_codigo = {c: Counter() for c in evaluables}
    no_verificadas = Counter()
    for fila in filas:
        detectadas = {r["codigo"] for r in fila["resultados"] if r["estado"] == "falla"}
        esperadas = set(fila["esperadas"]) & evaluables
        for r in fila["resultados"]:
            if r["estado"] == "no_verificado":
                no_verificadas[r["codigo"]] += 1
        for codigo in esperadas:
            por_codigo[codigo]["esperada"] += 1
            if codigo in detectadas:
                por_codigo[codigo]["detectada"] += 1
                vp += 1
            else:
                fn += 1
        for codigo in detectadas - esperadas:
            por_codigo[codigo]["falso_positivo"] += 1
            fp += 1
    return {
        "verdaderos_positivos": vp, "falsos_negativos": fn, "falsos_positivos": fp,
        "exhaustividad": round(vp / (vp + fn), 4) if vp + fn else None,
        "precision": round(vp / (vp + fp), 4) if vp + fp else None,
        "por_codigo": {c: dict(v) for c, v in por_codigo.items() if v},
        "no_verificadas": dict(no_verificadas),
        "consultas": sum(f["consultas"] for f in filas), "tokens": sum(f["tokens"] for f in filas),
        "segundos": round(sum(f["segundos"] for f in filas), 1),
    }


def _porcentaje(valor) -> str:
    return "—" if valor is None else f"{valor * 100:.1f} %".replace(".", ",")


def reporte(filas: list[dict], m: dict, con_agente: bool) -> str:
    lineas = [
        "# Validación del módulo de revisión de Bases", "",
        f"- **Fecha:** {datetime.now():%d/%m/%Y %H:%M}",
        f"- **Modo:** {'reglas fijas y agente revisor' if con_agente else 'solo reglas fijas'}",
        f"- **Convocatorias:** {len({f['convocatoria'] for f in filas})} · rondas revisadas: {len(filas)}", "",
        "## Resumen", "",
        "| Métrica | Valor |", "|---|---|",
        f"| Fallas sembradas detectadas (exhaustividad) | {_porcentaje(m['exhaustividad'])} "
        f"({m['verdaderos_positivos']} de {m['verdaderos_positivos'] + m['falsos_negativos']}) |",
        f"| Hallazgos que eran fallas reales (precisión) | {_porcentaje(m['precision'])} |",
        f"| Falsos positivos | {m['falsos_positivos']} |",
        f"| Verificaciones no verificadas | {sum(m['no_verificadas'].values())} |",
        f"| Consultas a Gemini | {m['consultas']} |",
        f"| Tokens | {m['tokens']:,} |".replace(",", " "),
        f"| Tiempo total | {m['segundos']:.0f} s |", "",
        "## Por convocatoria y ronda", "",
        "| Convocatoria | Ronda | Esperadas | Detectadas | No detectadas | De más |", "|---|---|---|---|---|---|",
    ]
    for f in filas:
        detectadas = sorted({r["codigo"] for r in f["resultados"] if r["estado"] == "falla"})
        esperadas = set(f["esperadas"])
        lineas.append(f"| {f['convocatoria']} | {f['ronda']} | {', '.join(f['esperadas']) or '—'} | "
                      f"{', '.join(detectadas) or '—'} | {', '.join(sorted(esperadas - set(detectadas))) or '—'} | "
                      f"{', '.join(sorted(set(detectadas) - esperadas)) or '—'} |")
    lineas += ["", "## Hallazgos de más (posibles falsos positivos)", ""]
    de_mas = [(f, r) for f in filas for r in f["resultados"]
              if r["estado"] == "falla" and r["codigo"] not in f["esperadas"]]
    if not de_mas:
        lineas.append("Ninguno.")
    for f, r in de_mas:
        metodo = "agente" if POR_CODIGO[r["codigo"]].metodo == AGENTE else "regla"
        lineas.append(f"- **{f['convocatoria']}, ronda {f['ronda']}, {r['codigo']}** ({metodo}): {r['descripcion']}")
    lineas += ["", "Un hallazgo de más no siempre es un error: puede ser una falla real que no se sembró a propósito. "
               "Cada uno se revisa a mano antes de reportarlo como falso positivo.", ""]
    return "\n".join(lineas)


def _tabla(doc, encabezados: list[str], filas: list[list[str]]) -> None:
    tabla = doc.add_table(rows=1, cols=len(encabezados))
    tabla.style = "Table Grid"
    for celda, texto in zip(tabla.rows[0].cells, encabezados):
        celda.text = texto
    for fila in filas:
        for celda, texto in zip(tabla.add_row().cells, fila):
            celda.text = str(texto)


def informe_word(filas: list[dict], m: dict, con_agente: bool, ruta: Path) -> None:
    """Informe de validación: qué se probó, con qué documentos, qué detectó el sistema y las métricas."""
    doc = Document()
    doc.add_heading("Validación del módulo de revisión de Bases", level=1)
    doc.add_paragraph(f"Fecha: {datetime.now():%d/%m/%Y %H:%M} · Modelo: {LLM_MODEL} · Modo: "
                      f"{'reglas fijas y agente revisor' if con_agente else 'solo reglas fijas'}")

    doc.add_heading("Qué se probó y cómo", level=2)
    doc.add_paragraph(
        "Se usaron convocatorias de demostración ficticias con fallas sembradas a propósito en el expediente, en las "
        "Bases y en el pliego de absolución. Para cada una se sabe de antemano qué verificaciones del catálogo deben "
        "fallar en cada ronda. El sistema revisó los documentos igual que en el portal del comité (ronda A con el "
        "expediente y las Bases convocadas; ronda B con el pliego y las Bases Integradas) y se comparó lo que detectó "
        "con lo esperado.")
    doc.add_paragraph("Exhaustividad = fallas sembradas detectadas ÷ fallas sembradas. Precisión = hallazgos que eran "
                      "fallas sembradas ÷ hallazgos del sistema. Un hallazgo de más no siempre es un error: puede ser una "
                      "falla real que no se sembró; por eso cada uno se lista para revisarlo.")

    doc.add_heading("Resultados globales", level=2)
    _tabla(doc, ["Métrica", "Valor"], [
        ["Exhaustividad (fallas sembradas detectadas)",
         f"{_porcentaje(m['exhaustividad'])} ({m['verdaderos_positivos']} de "
         f"{m['verdaderos_positivos'] + m['falsos_negativos']})"],
        ["Precisión", _porcentaje(m["precision"])], ["Hallazgos de más", m["falsos_positivos"]],
        ["Verificaciones no verificadas", sum(m["no_verificadas"].values())],
        ["Consultas a Gemini", m["consultas"]], ["Tokens", f"{m['tokens']:,}".replace(",", " ")],
        ["Tiempo total", f"{m['segundos']:.0f} s"],
    ])

    for carpeta in dict.fromkeys(f["convocatoria"] for f in filas):
        de_la_convocatoria = [f for f in filas if f["convocatoria"] == carpeta]
        doc.add_heading(de_la_convocatoria[0].get("concurso", carpeta), level=2)
        doc.add_paragraph(f"Fallas sembradas: {EXPEDIENTES[carpeta].descripcion}")
        if de_la_convocatoria[0].get("concurso_id"):
            doc.add_paragraph("En el aplicativo: portal del comité → Concursos en Curso → "
                              f"«{de_la_convocatoria[0]['concurso']}».")
        doc.add_paragraph("Documentos usados:")
        _tabla(doc, ["Ronda", "Archivo", "Tipo de documento", "Páginas"], [
            [f["ronda"], d["archivo"], TIPOS_DOCUMENTO[d["tipo"]], d["paginas"]]
            for f in de_la_convocatoria for d in f.get("documentos", [])])
        for f in de_la_convocatoria:
            detectadas = sorted({r["codigo"] for r in f["resultados"] if r["estado"] == "falla"})
            esperadas = set(f["esperadas"])
            doc.add_heading(f"Ronda {f['ronda']}", level=3)
            _tabla(doc, ["Esperadas", "Detectadas", "No detectadas", "De más"], [[
                ", ".join(f["esperadas"]) or "—", ", ".join(detectadas) or "—",
                ", ".join(sorted(esperadas - set(detectadas))) or "—",
                ", ".join(sorted(set(detectadas) - esperadas)) or "—"]])
            hallazgos = [r for r in f["resultados"] if r["estado"] != "cumple"]
            if hallazgos:
                doc.add_paragraph("Lo que reportó el sistema:")
                _tabla(doc, ["Verificación", "Resultado", "Descripción y pasos del agente"], [
                    [f"{r['codigo']} · {POR_CODIGO[r['codigo']].nombre}",
                     "Falla" if r["estado"] == "falla" else "No verificado",
                     r["descripcion"] + ("\nPasos: " + " → ".join(p["resumen"] for p in r["pasos"]) if r.get("pasos") else "")]
                    for r in hallazgos])
            doc.add_paragraph(f"Consultas a Gemini: {f['consultas']} · tiempo: {f['segundos']:.0f} s.")
    doc.save(ruta)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--sin-agente", action="store_true")
    parser.add_argument("--en-memoria", action="store_true", help="no escribe en Supabase")
    parser.add_argument("--convocatorias", default="", help="carpetas separadas por comas (por defecto, todas)")
    args = parser.parse_args()
    carpetas = [c for c in args.convocatorias.split(",") if c] or list(EXPEDIENTES)
    salida = CARPETA_RESULTADOS / f"{datetime.now():%Y-%m-%d_%H%M}_revision{'_reglas' if args.sin_agente else ''}"
    salida.mkdir(parents=True, exist_ok=True)
    filas = []
    for carpeta in carpetas:
        print(f"Revisando {carpeta}...", flush=True)
        filas += revisar_convocatoria(carpeta, not args.sin_agente, args.en_memoria)
        # Se guarda lo avanzado después de cada convocatoria, por si la ejecución se corta.
        (salida / "resultados.json").write_text(json.dumps({"filas": filas}, ensure_ascii=False, indent=1),
                                                encoding="utf-8")
    m = metricas(filas, not args.sin_agente)
    (salida / "resultados.json").write_text(json.dumps({"filas": filas, "metricas": m}, ensure_ascii=False, indent=1),
                                            encoding="utf-8")
    (salida / "REPORTE_REVISION.md").write_text(reporte(filas, m, not args.sin_agente), encoding="utf-8")
    informe_word(filas, m, not args.sin_agente, salida / "VALIDACION_REVISION.docx")
    print(f"Exhaustividad {_porcentaje(m['exhaustividad'])} · precisión {_porcentaje(m['precision'])} · "
          f"{m['consultas']} consultas. Reporte: {salida / 'VALIDACION_REVISION.docx'}")


if __name__ == "__main__":
    main()
