from docx import Document

from reports.acta import generar_acta

CONCURSO = {"nombre": "Supervisión vial"}
POSTOR = {"razon_social": "Consultora SAC", "ruc": "20131312955"}
FILAS = [
    {"criterio_descripcion": "Declaración jurada", "etapa": "admision", "veredicto": "Cumple",
     "justificacion": "Presenta el Anexo 2.", "citas": [{"documento": "propuesta", "pagina": 2, "extracto": "Anexo 2"}]},
    {"criterio_descripcion": "Experiencia del postor", "etapa": "evaluacion", "veredicto": "Requiere revisión",
     "puntaje": 20, "puntaje_maximo": 60, "justificacion": "Un contrato es ilegible.",
     "veredicto_comite": "Evaluado", "puntaje_comite": 40, "justificacion_comite": "Se verificó el contrato original."},
]


def _texto(documento) -> str:
    celdas = [c.text for t in documento.tables for fila in t.rows for c in fila.cells]
    return "\n".join([p.text for p in documento.paragraphs] + celdas)


def test_acta_revisada_muestra_ia_y_comite_por_etapa():
    dictamen = {"veredicto": "Requiere revisión", "veredicto_comite": "Califica", "puntaje": 20, "puntaje_comite": 40,
                "puntaje_maximo": 60, "justificacion": "Conciliación.", "revisado_por": "Ing. Ana Torres",
                "revisado_en": "2026-10-12T15:30:00+00:00", "justificacion_comite": "Sin observaciones."}

    texto = _texto(Document(generar_acta(CONCURSO, POSTOR, FILAS, dictamen)))

    assert "Resultado final: Califica" in texto
    assert "Puntaje técnico: 40 de 60" in texto
    assert "Ing. Ana Torres, 12/10/2026 15:30" in texto and "Propuesta original de la IA: Requiere revisión" in texto
    assert "20 de 60 puntos (requiere revisión)" in texto and "40 de 60 puntos" in texto
    assert "Propuesta Técnica, página 2: Anexo 2" in texto
    assert "Comité: Se verificó el contrato original." in texto


def test_acta_sin_revisar_lo_indica():
    dictamen = {"veredicto": "Califica", "puntaje": 20, "puntaje_maximo": 60, "justificacion": "ok"}
    texto = _texto(Document(generar_acta(CONCURSO, POSTOR, FILAS, dictamen)))
    assert "pendiente de revisión por el comité" in texto
