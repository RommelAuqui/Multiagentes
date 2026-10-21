from validacion.revision import metricas, reporte


def _fila(convocatoria, ronda, esperadas, fallas, no_verificadas=()):
    resultados = [{"codigo": c, "estado": "falla", "descripcion": f"Falla {c}"} for c in fallas]
    resultados += [{"codigo": c, "estado": "no_verificado", "descripcion": "Falta"} for c in no_verificadas]
    return {"convocatoria": convocatoria, "ronda": ronda, "esperadas": esperadas, "resultados": resultados,
            "segundos": 10.0, "consultas": 5, "tokens": 1000}


def test_metricas_de_deteccion():
    filas = [
        _fila("conv4", "A", ["COH-02", "COH-06"], ["COH-02", "PRO-05"]),  # detecta 1 de 2 y un hallazgo de más
        _fila("conv1", "A", [], [], ["FOR-01"]),
    ]
    m = metricas(filas, con_agente=True)

    assert (m["verdaderos_positivos"], m["falsos_negativos"], m["falsos_positivos"]) == (1, 1, 1)
    assert m["exhaustividad"] == 0.5 and m["precision"] == 0.5
    assert m["por_codigo"]["COH-06"] == {"esperada": 1} and m["por_codigo"]["PRO-05"] == {"falso_positivo": 1}
    assert m["no_verificadas"] == {"FOR-01": 1} and m["consultas"] == 10

    texto = reporte(filas, m, con_agente=True)
    assert "| conv4 | A | COH-02, COH-06 | COH-02, PRO-05 | COH-06 | PRO-05 |" in texto
    assert "**conv4, ronda A, PRO-05** (agente): Falla PRO-05" in texto


def test_el_informe_word_lista_los_documentos_y_los_resultados(tmp_path):
    from docx import Document

    from validacion.revision import informe_word

    fila = _fila("convocatoria_4_supervision_canal_de_riego", "A", ["COH-02"], ["COH-02"])
    fila.update({"concurso": "VALIDACIÓN · CP-04-2026-PERVS/CS · supervision canal de riego", "concurso_id": "c1",
                 "documentos": [{"tipo": "certificacion", "archivo": "A2 Certificación presupuestal.pdf",
                                 "version": 1, "paginas": 1}]})
    ruta = tmp_path / "informe.docx"
    informe_word([fila], metricas([fila], con_agente=True), True, ruta)

    doc = Document(ruta)
    textos = [p.text for p in doc.paragraphs]
    celdas = [c.text for t in doc.tables for f in t.rows for c in f.cells]
    assert "VALIDACIÓN · CP-04-2026-PERVS/CS · supervision canal de riego" in textos
    assert "A2 Certificación presupuestal.pdf" in celdas and "100,0 % (1 de 1)" in celdas
    assert any("Concursos en Curso" in t for t in textos)


def test_sin_agente_solo_se_cuentan_las_reglas_fijas():
    # COH-06 la revisa el agente: sin él no se espera que se detecte.
    m = metricas([_fila("conv4", "A", ["COH-02", "COH-06"], ["COH-02"])], con_agente=False)
    assert m["exhaustividad"] == 1.0 and m["falsos_negativos"] == 0
