import pytest

from rag.pdf_extract import ResultadoExtraccion
from revision import servicio
from revision.reglas import FALLA, Resultado

PAGINAS_BASES = [
    {"pagina": 1, "texto": "CAPÍTULO I GENERALIDADES CAPÍTULO II DEL PROCEDIMIENTO DE SELECCIÓN"},
    {"pagina": 2, "texto": "CAPÍTULO III REQUERIMIENTO CAPÍTULO IV EVALUACIÓN DE OFERTAS CAPÍTULO V PROFORMA"},
]
FICHA_BASES = {
    "cuantia": 480000.0, "fuente_financiamiento": "Recursos determinados", "fecha_aprobacion_expediente": "2026-09-28",
    "difusion_requerimiento": False, "elevacion_oece": True, "plazo_dias": 210.0, "monto_experiencia_postor": 480000.0,
    "puntaje_minimo": 70.0, "puntaje_total": 100.0,
    "factores": [{"nombre": "Experiencia", "puntaje": 60.0}, {"nombre": "Metodología", "puntaje": 40.0}],
    "cronograma": {"convocatoria": "2026-10-05", "consultas_fin": "2026-10-16"},
}


@pytest.fixture
def fichas(monkeypatch):
    """Simula la lectura del PDF y la extracción: el «contenido» del archivo dice qué ficha devolver."""
    por_contenido = {
        b"requerimiento": {"plazo_dias": 210.0},
        b"ccp-insuficiente": {"monto": 400000.0, "fuente_financiamiento": "Recursos determinados"},
        b"ccp-corregida": {"monto": 480000.0, "fuente_financiamiento": "Recursos determinados"},
        b"aprobacion": {"fecha_aprobacion": "2026-09-28", "cuantia": 480000.0},
        b"bases": FICHA_BASES,
    }
    consultas = []

    def leer(contenido, contexto):
        paginas = PAGINAS_BASES if contenido == b"bases" else [{"pagina": 1, "texto": contenido.decode()}]
        return ResultadoExtraccion(total_paginas=len(paginas), paginas=paginas)

    def extraer(tipo, texto, contexto):
        consultas.append(tipo)
        if tipo == "bases":
            return dict(FICHA_BASES)
        return next(dict(f) for clave, f in por_contenido.items() if clave.decode() in texto)

    monkeypatch.setattr(servicio, "extraer_paginas", leer)
    # El agente revisor tiene sus propias pruebas; aquí no encuentra nada.
    monkeypatch.setattr(servicio.agente_revisor, "revisar", lambda _ronda, _docs, _contexto: [])
    monkeypatch.setattr(servicio, "extraer_ficha", extraer)
    return consultas


@pytest.fixture
def concurso(db_falsa):
    return db_falsa.table("concursos").insert({"nombre": "Prueba", "estado": "revision_bases",
                                               "ronda_revision": "A"}).execute().data[0]


def _subir_ronda_a(concurso_id, certificacion=b"ccp-insuficiente"):
    for tipo, contenido in (("requerimiento", b"requerimiento"), ("certificacion", certificacion),
                            ("aprobacion_expediente", b"aprobacion"), ("bases", b"bases")):
        servicio.subir_documento(concurso_id, tipo, f"{tipo}.pdf", contenido)


def _por_codigo(concurso_id):
    return {h["codigo"]: h for h in servicio.hallazgos(concurso_id, "A")}


def test_flujo_completo_de_la_ronda_a(concurso, fichas, db_falsa):
    cid = concurso["id"]
    assert servicio.puede_continuar(cid, "A")[0] is False
    _subir_ronda_a(cid)
    assert fichas == ["requerimiento", "certificacion", "aprobacion_expediente", "bases"]
    assert "Hay documentos nuevos sin analizar" in servicio.puede_continuar(cid, "A")[1]

    primera = servicio.revisar(cid, "A")
    hallazgos = _por_codigo(cid)
    # La certificación no alcanza y faltan los documentos opcionales del expediente.
    assert set(hallazgos) == {"COH-02", "EXP-04", "EXP-05", "EXP-06"}
    assert hallazgos["COH-02"]["gravedad"] == "Crítico" and hallazgos["COH-02"]["estado"] == "pendiente"
    assert hallazgos["COH-02"]["detectado_en"] == primera["id"]
    assert "Quedan 4 hallazgo(s) sin responder" in servicio.puede_continuar(cid, "A")[1]

    with pytest.raises(servicio.RespuestaInvalida):
        servicio.mantener(hallazgos["EXP-06"]["id"], "  ", "Ana Pérez")
    servicio.mantener(hallazgos["EXP-06"]["id"], "La designación consta en la resolución del expediente.", "Ana Pérez")

    # Se corrige la certificación: se sube la versión 2 y se vuelve a analizar.
    servicio.subir_documento(cid, "certificacion", "certificacion_v2.pdf", b"ccp-corregida")
    assert servicio.documentos_vigentes(cid)["certificacion"]["version"] == 2
    segunda = servicio.revisar(cid, "A")
    hallazgos = _por_codigo(cid)
    assert hallazgos["COH-02"]["estado"] == "corregido" and hallazgos["COH-02"]["resuelto_en"] == segunda["id"]
    assert hallazgos["EXP-06"]["estado"] == "se_mantiene" and hallazgos["EXP-06"]["revisado_por"] == "Ana Pérez"

    for codigo in ("EXP-04", "EXP-05"):
        servicio.mantener(hallazgos[codigo]["id"], "Contratación rutinaria: no se elaboró ese formato.", "Ana Pérez")
    assert servicio.puede_continuar(cid, "A") == (True, "")
    # El historial queda: dos revisiones y las dos versiones de la certificación.
    assert len(db_falsa.filas("revisiones_bases")) == 2
    assert [d["version"] for d in db_falsa.filas("documentos_expediente") if d["tipo"] == "certificacion"] == [1, 2]


def test_un_hallazgo_corregido_que_reaparece_vuelve_a_quedar_pendiente():
    anterior = {"id": "h1", "codigo": "COH-02", "clave": "", "estado": "corregido", "justificacion_comite": "x"}
    nuevos, cambios = servicio.conciliar([anterior], [Resultado("COH-02", FALLA, "Otra vez no alcanza")], "r3")

    assert nuevos == []
    [(hallazgo_id, datos)] = cambios
    assert hallazgo_id == "h1" and datos["estado"] == "pendiente" and datos["justificacion_comite"] is None


def test_una_consulta_respondida_desaparece_y_su_hallazgo_queda_corregido():
    anteriores = [{"id": "h1", "codigo": "ABS-01", "clave": "2", "estado": "pendiente"},
                  {"id": "h2", "codigo": "COH-02", "clave": "", "estado": "pendiente"}]
    resultados = [Resultado("ABS-01", "cumple", "Todas tienen respuesta.")]

    nuevos, cambios = servicio.conciliar(anteriores, resultados, "r2")

    # ABS-01 se revisó y la consulta 2 ya no falla; COH-02 no se revisó en esta corrida y no se toca.
    assert nuevos == [] and [(i, d["estado"]) for i, d in cambios] == [("h1", "corregido")]


def test_el_informe_y_el_acta_dejan_constancia_de_lo_mantenido(concurso, fichas):
    from docx import Document

    from reports.acta import generar_acta

    cid = concurso["id"]
    _subir_ronda_a(cid)
    servicio.revisar(cid, "A")
    for hallazgo in servicio.hallazgos(cid, "A"):
        servicio.mantener(hallazgo["id"], "Se gestiona la ampliación presupuestal.", "Ana Pérez")

    texto = "\n".join(p.text for p in Document(servicio.informe(concurso)).paragraphs)
    tablas = "\n".join(c.text for t in Document(servicio.informe(concurso)).tables for f in t.rows for c in f.cells)
    assert "4 hallazgo(s): 0 corregido(s) y 4 mantenido(s)" in texto
    assert "mantener 1 hallazgo(s) crítico(s)" in texto
    assert "COH-02 · La certificación presupuestal cubre la cuantía" in tablas and "Firma: Ana Pérez" in tablas

    lineas = servicio.resumen_para_el_acta(cid)
    assert any("Hallazgo crítico mantenido COH-02" in linea for linea in lineas)
    acta = Document(generar_acta(concurso, {"razon_social": "Consultora", "ruc": "20100000001"}, [],
                                 {"veredicto": "Califica", "justificacion": ""}, lineas))
    assert "Control previo de las Bases" in [p.text for p in acta.paragraphs]


def test_un_concurso_sin_revision_de_bases_no_agrega_nada_al_acta(concurso):
    assert servicio.resumen_para_el_acta(concurso["id"]) is None
