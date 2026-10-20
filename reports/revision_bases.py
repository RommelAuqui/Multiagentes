"""Informe de revisión de las Bases (Word) y su resumen para el acta de evaluación.

Deja constancia de qué se revisó, qué se encontró en cada ronda, qué se corrigió y qué decidió mantener el
comité, con su justificación y su firma.
"""
from datetime import datetime
from io import BytesIO

from docx import Document

from consumo import ZONA_HORARIA
from revision.catalogo import CRITICO, POR_CODIGO, TIPOS_DOCUMENTO

RONDAS = {"A": "Ronda A: antes de convocar", "B": "Ronda B: antes de recibir ofertas"}
ESTADOS = {"pendiente": "Pendiente de respuesta", "se_mantiene": "Se mantiene", "corregido": "Corregido"}
RESULTADOS = {"falla": "Falla", "no_verificado": "No verificado", "cumple": "Cumple"}


def _fecha(valor: str | None) -> str:
    try:
        return datetime.fromisoformat(valor).astimezone(ZONA_HORARIA).strftime("%d/%m/%Y %H:%M")
    except (TypeError, ValueError):
        return "—"


def _evidencia(hallazgo: dict) -> str:
    partes = [hallazgo["descripcion"]]
    if citas := [c for c in hallazgo.get("citas") or [] if c.get("documento")]:
        partes.append("Dónde: " + "; ".join(c["documento"] + (f", p. {c['pagina']}" if c.get("pagina") else "")
                                             for c in citas))
    if hallazgo.get("pasos"):
        partes.append("Proceso del agente revisor: " + " → ".join(p["resumen"] for p in hallazgo["pasos"]))
    return "\n".join(partes)


def _respuesta(hallazgo: dict) -> str:
    if hallazgo["estado"] == "se_mantiene":
        return (f"Se mantiene. {hallazgo.get('justificacion_comite') or ''}\n"
                f"Firma: {hallazgo.get('revisado_por') or '—'}, {_fecha(hallazgo.get('revisado_en'))}")
    if hallazgo["estado"] == "corregido":
        return "Corregido: lo confirmó una revisión posterior del sistema."
    return "Pendiente de respuesta."


def resumen_para_acta(hallazgos: list[dict]) -> list[str]:
    """Párrafos sobre la revisión de Bases para incluir en el acta de evaluación de cada postor."""
    if not hallazgos:
        return ["Las Bases pasaron la revisión de control previo sin hallazgos."]
    corregidos = sum(h["estado"] == "corregido" for h in hallazgos)
    mantenidos = [h for h in hallazgos if h["estado"] == "se_mantiene"]
    lineas = [f"La revisión de control previo de las Bases registró {len(hallazgos)} hallazgo(s): {corregidos} "
              f"corregido(s) y {len(mantenidos)} mantenido(s) por decisión del comité."]
    for h in mantenidos:
        gravedad = "crítico" if h["gravedad"] == CRITICO else h["gravedad"].lower()
        lineas.append(f"Hallazgo {gravedad} mantenido {h['codigo']} ({POR_CODIGO[h['codigo']].nombre}): "
                      f"{h.get('justificacion_comite') or ''} — {h.get('revisado_por') or '—'}, "
                      f"{_fecha(h.get('revisado_en'))}.")
    return lineas


def generar_informe(concurso: dict, hallazgos: list[dict], revisiones: list[dict], documentos: list[dict]) -> BytesIO:
    doc = Document()
    doc.add_heading("Informe de revisión de las Bases", level=1)
    doc.add_paragraph(f"Concurso: {concurso['nombre']}")
    doc.add_paragraph(f"Fecha del informe: {datetime.now(ZONA_HORARIA):%d/%m/%Y %H:%M}")

    doc.add_heading("Resumen", level=2)
    for linea in resumen_para_acta(hallazgos):
        doc.add_paragraph(linea)
    criticos = [h for h in hallazgos if h["gravedad"] == CRITICO and h["estado"] == "se_mantiene"]
    if criticos:
        doc.add_paragraph(f"Atención: el comité decidió mantener {len(criticos)} hallazgo(s) crítico(s).")
    doc.add_paragraph(f"Revisiones realizadas: {len(revisiones)} (cada «Analizar» queda registrado con los documentos "
                      "que se usaron).")

    doc.add_heading("Documentos revisados", level=2)
    tabla = doc.add_table(rows=1, cols=4)
    tabla.style = "Table Grid"
    for celda, texto in zip(tabla.rows[0].cells, ("Tipo de documento", "Archivo", "Versión", "Subido")):
        celda.text = texto
    for d in sorted(documentos, key=lambda d: (list(TIPOS_DOCUMENTO).index(d["tipo"]), d["version"])):
        fila = tabla.add_row().cells
        fila[0].text, fila[1].text = TIPOS_DOCUMENTO[d["tipo"]], d["nombre"]
        fila[2].text, fila[3].text = str(d["version"]), _fecha(d.get("created_at"))

    for ronda, titulo in RONDAS.items():
        de_la_ronda = [h for h in hallazgos if h["ronda"] == ronda]
        if not de_la_ronda and not any(r["ronda"] == ronda for r in revisiones):
            continue
        doc.add_heading(titulo, level=2)
        if not de_la_ronda:
            doc.add_paragraph("Sin hallazgos.")
            continue
        tabla = doc.add_table(rows=1, cols=4)
        tabla.style = "Table Grid"
        for celda, texto in zip(tabla.rows[0].cells, ("Verificación", "Resultado", "Hallazgo y evidencia",
                                                       "Respuesta del comité")):
            celda.text = texto
        for h in de_la_ronda:
            verificacion = POR_CODIGO[h["codigo"]]
            fila = tabla.add_row().cells
            fila[0].text = f"{h['codigo']} · {verificacion.nombre}\n{h['gravedad']}\nRespaldo: {verificacion.respaldo}"
            fila[1].text = f"{RESULTADOS[h['resultado']]}\n{ESTADOS[h['estado']]}"
            fila[2].text = _evidencia(h)
            fila[3].text = _respuesta(h)

    salida = BytesIO()
    doc.save(salida)
    salida.seek(0)
    return salida
