from datetime import datetime
from io import BytesIO

from docx import Document

from agents.reglas import ETAPAS, REQUIERE_REVISION, puntaje_efectivo, veredicto_efectivo

# Ofertas de la versión anterior: dos documentos fijos en lugar de archivos con nombre propio.
NOMBRES_DOCUMENTO = {"propuesta": "Propuesta Técnica", "cv": "CV Documentado"}
ENCABEZADOS = ("Criterio", "Propuesta de la IA", "Decisión del comité", "Sustento y citas")


def nombre_legible(documento: str | None) -> str | None:
    return NOMBRES_DOCUMENTO.get(documento, documento)


def _texto_cita(cita: dict) -> str:
    documento = nombre_legible(cita.get("documento"))
    ubicacion = f"{documento}, página {cita.get('pagina')}" if documento else f"Página {cita.get('pagina')}"
    return f"{ubicacion}: {cita.get('extracto', '')}"


def _fecha(valor: str) -> str:
    try:
        return datetime.fromisoformat(valor).strftime("%d/%m/%Y %H:%M")
    except (TypeError, ValueError):
        return str(valor)


def _numero(valor) -> str:
    return f"{float(valor):g}"


def _decision(fila: dict, del_comite: bool) -> str:
    veredicto = fila.get("veredicto_comite") if del_comite else fila["veredicto"]
    if not veredicto:
        return "—"
    puntaje = fila.get("puntaje_comite") if del_comite else fila.get("puntaje")
    if fila.get("etapa") != "evaluacion" or puntaje is None:
        return veredicto
    texto = f"{_numero(puntaje)} de {_numero(fila.get('puntaje_maximo') or 0)} puntos"
    return f"{texto} ({REQUIERE_REVISION.lower()})" if veredicto == REQUIERE_REVISION else texto


def _sustento(fila: dict) -> str:
    partes = [fila.get("justificacion") or ""]
    if fila.get("justificacion_comite"):
        partes.append(f"Comité: {fila['justificacion_comite']}")
    if fila.get("citas"):
        partes.append("Citas: " + "; ".join(_texto_cita(c) for c in fila["citas"]))
    if fila.get("pasos"):
        partes.append("Proceso del agente: " + " → ".join(p["resumen"] for p in fila["pasos"]))
    return "\n".join(p for p in partes if p)


def generar_acta(concurso: dict, postor: dict, filas: list[dict], dictamen: dict) -> BytesIO:
    doc = Document()
    doc.add_heading("Acta de Dictamen Técnico", level=1)
    doc.add_paragraph(f"Concurso: {concurso['nombre']}")
    doc.add_paragraph(f"Postor: {postor['razon_social']} (RUC {postor['ruc']})")

    doc.add_heading("Resultado", level=2)
    doc.add_paragraph(f"Resultado final: {veredicto_efectivo(dictamen)}")
    if dictamen.get("puntaje_maximo"):
        doc.add_paragraph(
            f"Puntaje técnico: {_numero(puntaje_efectivo(dictamen) or 0)} de {_numero(dictamen['puntaje_maximo'])}"
        )
    if dictamen.get("revisado_en"):
        doc.add_paragraph(
            f"Revisado por el comité de selección: {dictamen['revisado_por']}, {_fecha(dictamen['revisado_en'])}. "
            f"Propuesta original de la IA: {dictamen['veredicto']}."
        )
        if dictamen.get("justificacion_comite"):
            doc.add_paragraph(f"Observaciones del comité: {dictamen['justificacion_comite']}")
    else:
        doc.add_paragraph("Propuesta de la IA pendiente de revisión por el comité de selección.")

    doc.add_heading("Conciliación del Agente Coordinador", level=2)
    doc.add_paragraph(dictamen.get("justificacion") or "")
    for paso in dictamen.get("pasos") or []:
        doc.add_paragraph(paso["resumen"])

    for etapa, nombre in ETAPAS.items():
        de_la_etapa = [f for f in filas if f.get("etapa") == etapa]
        if not de_la_etapa:
            continue
        doc.add_heading(nombre, level=2)
        tabla = doc.add_table(rows=1, cols=len(ENCABEZADOS))
        tabla.style = "Table Grid"
        for celda, encabezado in zip(tabla.rows[0].cells, ENCABEZADOS):
            celda.text = encabezado
        for fila in de_la_etapa:
            celdas = tabla.add_row().cells
            celdas[0].text = fila["criterio_descripcion"]
            celdas[1].text = _decision(fila, del_comite=False)
            celdas[2].text = _decision(fila, del_comite=True)
            celdas[3].text = _sustento(fila)

    doc.add_paragraph(
        "Documento generado con apoyo de un sistema multiagente de inteligencia artificial. "
        "La decisión corresponde al comité de selección."
    )

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
